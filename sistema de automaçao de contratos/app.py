from openpyxl import load_workbook
from doxc import Document
from datetime import datetime

planilhaForncedor = load_workbook('./fornercedores.xlsx')
paginaForncedor = planilhaForncedor['Sheet1']

for linha in paginaForncedor.inter_rows(min_row=2, values_only=True):
    nomeEmpresa,endereco, cidade, estado, cep, telefone, email, setor = linha

    arquivoWord = Document()
    arquivoWord.add_heading('Contrato de prestação',0)

    texto_contrato = f"""
    Este contrato de prestação de serviços é feito entre {nomeEmpresa}, com endereço em {endereco}, 
    {cidade}, {estado}, CEP {cep}, doravante denominado FORNECEDOR, e a empresa CONTRATANTE.

    Pelo presente instrumento particular, as partes têm, entre si, justo e acordado o seguinte:

    1. OBJETO DO CONTRATO
    O FORNECEDOR compromete-se a fornecer à CONTRATANTE os serviços/material de acordo com as especificações acordadas, respeitando os padrões de qualidade e os prazos estipulados.

    2. PRAZO
    Este contrato tem prazo de vigência de 12 (doze) meses, iniciando-se na data de sua assinatura, podendo ser renovado conforme acordo entre as partes.

    3. VALOR E FORMA DE PAGAMENTO
    O valor dos serviços prestados será acordado conforme as demandas da CONTRATANTE e a capacidade de entrega do FORNECEDOR. Os pagamentos serão realizados mensalmente, mediante apresentação de nota fiscal.

    4. CONFIDENCIALIDADE
    Todas as informações trocadas entre as partes durante a vigência deste contrato serão tratadas como confidenciais.

    Para firmeza e como prova de assim haverem justo e contratado, as partes assinam o presente contrato em duas vias de igual teor e forma.

    FORNECEDOR: {nomeEmpresa}
    E-mail: {email}

    CONTRATANTE: Contratante da Silva
    E-mail: contratabte@gamil.com

    São Paulo, {datetime.now().strftime('%d/%m/%Y')}"""
    
    arquivoWord.add_paragraph(texto_contrato)

    arquivoWord.save(f'./contratos_{nomeEmpresa}.docx')